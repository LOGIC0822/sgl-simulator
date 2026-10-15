#!/usr/bin/env python3
"""Summarize AIC-vs-real one-batch forward measurements.

The real measurements are produced by upstream ``sglang.benchmark.one_batch``
with CUDA Graph enabled.  AIC predicts an eager static operator composition,
so this report deliberately labels the comparison as deployment-step error,
not as a CUDA-Graph replay measurement by AIC.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
from pathlib import Path
from typing import Any


PHASES = (
    ("prefill", "prefill_latency"),
    ("decode", "median_decode_latency"),
)


def read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object in {path}")
    return value


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        value = json.loads(line)
        if not isinstance(value, dict):
            raise ValueError(f"Expected JSON object at {path}:{line_number}")
        rows.append(value)
    return rows


def numeric(value: Any, label: str) -> float:
    if not isinstance(value, (float, int)) or isinstance(value, bool):
        raise ValueError(f"{label} is not numeric: {value!r}")
    value = float(value)
    if not math.isfinite(value) or value <= 0:
        raise ValueError(f"{label} must be finite and positive: {value!r}")
    return value


def fmt(value: float | None) -> str:
    return "—" if value is None else f"{value:.3f}"


def write_docx(path: Path, rows: list[dict[str, Any]], metadata: dict[str, str]) -> None:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt
    except ImportError:
        return

    document = Document()

    def set_simsun_rfonts(rpr: Any) -> None:
        fonts = rpr.get_or_add_rFonts()
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "SimSun")

    def set_simsun_run(run: Any) -> None:
        run.font.name = "SimSun"
        set_simsun_rfonts(run._element.get_or_add_rPr())

    def set_simsun_style(style: Any) -> None:
        style.font.name = "SimSun"
        set_simsun_rfonts(style._element.get_or_add_rPr())

    def walk_paragraphs(container: Any):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk_paragraphs(cell)

    def enforce_simsun() -> None:
        for style in document.styles:
            try:
                set_simsun_style(style)
            except AttributeError:
                continue
        containers = [document]
        for section in document.sections:
            containers.extend((section.header, section.footer))
        for container in containers:
            for paragraph in walk_paragraphs(container):
                for run in paragraph.runs:
                    set_simsun_run(run)

    document.styles["Normal"].font.name = "SimSun"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading("AIC 与真机单步 Forward 对照", 0)
    document.add_paragraph(
        "真实值来自 A100 上的 SGLang ModelRunner 单 batch 测量；"
        "AIC 为 Silicon 静态算子组合预测，不物理回放 CUDA Graph。"
    )
    document.add_paragraph(
        f"模型：{metadata.get('models', '未记录')}；输入长度："
        f"{metadata.get('input_len', '未记录')}；每模型真实重复："
        f"{metadata.get('repeats', '未记录')}。"
    )
    document.add_paragraph(
        "静态显存比例：默认 "
        f"{metadata.get('mem_fraction_static_default', '未记录')}；逐模型覆盖 "
        f"{metadata.get('mem_fraction_static_overrides', '无')}。"
    )
    document.add_paragraph(
        "内核路径：" + metadata.get("kernel_execution_path", "未记录")
    )
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    titles = ("模型", "阶段", "真实均值 ms", "真实标准差", "AIC ms", "带符号偏差", "MAPE", "样本")
    for cell, title in zip(table.rows[0].cells, titles):
        cell.text = title
    for row in rows:
        cells = table.add_row().cells
        values = (
            row["model"],
            row["phase"],
            fmt(row["real_mean_ms"]),
            fmt(row["real_stddev_ms"]),
            fmt(row["aic_predicted_ms"]),
            f"{fmt(row['signed_error_pct'])}%",
            f"{fmt(row['absolute_error_pct'])}%",
            str(row["real_samples"]),
        )
        for cell, value in zip(cells, values):
            cell.text = value
    document.add_heading("阅读边界", level=1)
    document.add_paragraph(
        "该结果用于判断 AIC 静态 forward 估算与实际部署单步之间的差异；"
        "它不包含 HTTP、排队或多请求调度时间，也不能将 CUDA Graph 的影响误写为 AIC 已模拟。"
    )
    enforce_simsun()
    document.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-root", type=Path, required=True)
    args = parser.parse_args()
    root = args.results_root.resolve()
    metadata_path = root / "run_metadata.txt"
    metadata: dict[str, str] = {}
    if metadata_path.is_file():
        for line in metadata_path.read_text(encoding="utf-8").splitlines():
            if "=" in line:
                key, value = line.split("=", 1)
                metadata[key] = value

    rows: list[dict[str, Any]] = []
    missing: list[str] = []
    for model_dir in sorted((root / "models").glob("*")):
        if not model_dir.is_dir():
            continue
        real_path = model_dir / "real" / "one_batch.jsonl"
        prediction_path = model_dir / "aic_prediction.json"
        if not real_path.is_file() or not prediction_path.is_file():
            missing.append(model_dir.name)
            continue
        real = read_jsonl(real_path)
        predicted = read_json(prediction_path)
        for phase, real_key in PHASES:
            samples = [numeric(item.get(real_key), f"{real_path}:{real_key}") * 1e3 for item in real]
            if not samples:
                raise ValueError(f"No {phase} samples in {real_path}")
            predicted_ms = numeric(
                predicted.get("phases", {}).get(phase, {}).get("predicted_ms"),
                f"{prediction_path}:{phase}.predicted_ms",
            )
            real_mean = statistics.mean(samples)
            signed_error = (predicted_ms - real_mean) / real_mean * 100.0
            rows.append(
                {
                    "model": model_dir.name,
                    "phase": phase,
                    "real_samples": len(samples),
                    "real_mean_ms": real_mean,
                    "real_stddev_ms": statistics.stdev(samples) if len(samples) > 1 else 0.0,
                    "real_min_ms": min(samples),
                    "real_max_ms": max(samples),
                    "aic_predicted_ms": predicted_ms,
                    "signed_error_pct": signed_error,
                    "absolute_error_pct": abs(signed_error),
                }
            )

    if missing:
        raise SystemExit("Missing completed results for: " + ", ".join(missing))
    if not rows:
        raise SystemExit(f"No model results found under {root / 'models'}")

    analysis = root / "analysis"
    analysis.mkdir(parents=True, exist_ok=True)
    csv_path = analysis / "single_step_comparison.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    lines = [
        "# AIC Silicon 与 A100 真机单步 Forward 对照",
        "",
        "## 测量定义",
        "",
        "- 真机：完整 checkpoint 的 SGLang `ModelRunner`，"
        f"`batch=1`、输入 {metadata.get('input_len', '未记录')} token、输出 "
        f"{metadata.get('output_len', '未记录')} token；每个模型同形状重复 "
        f"{metadata.get('repeats', '未记录')} 次。prefill 为一次 extend forward，"
        "decode 为首次生成后的一个 decode forward。",
        "- 真机保持 `triton` attention、decode `full` CUDA Graph、prefill `tc_piecewise` CUDA Graph；日志和 JSONL 位于 `models/<model>/real/`。",
        "- 静态显存比例：默认 `"
        + metadata.get("mem_fraction_static_default", "未记录")
        + "`；逐模型覆盖 `"
        + metadata.get("mem_fraction_static_overrides", "无")
        + "`。每个模型的实际值记录在 `models/<model>/real/launch-settings.txt`。",
        "- AIC：同一模型、同一形状、同一 A100 PCIe Silicon 表的静态算子组合预测。AIC 不物理回放 CUDA Graph，因此下表是“部署单步误差”，不是 AIC 已模拟 CUDA Graph 的证明。",
        "- 误差公式：`(AIC − 真机) / 真机 × 100%`；MAPE 为其绝对值。该对照不包含 HTTP、排队和多请求 scheduler 时间。",
        "",
        "## 结果",
        "",
        "| 模型 | 阶段 | 真机均值 ms | 标准差 ms | AIC 预测 ms | 带符号偏差 | MAPE | 真机样本 |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    for row in rows:
        lines.append(
            "| {model} | {phase} | {real_mean_ms:.3f} | {real_stddev_ms:.3f} | "
            "{aic_predicted_ms:.3f} | {signed_error_pct:+.2f}% | "
            "{absolute_error_pct:.2f}% | {real_samples} |".format(**row)
        )
    lines += [
        "",
        "## 复现信息",
        "",
        f"- PR 源码：`{metadata.get('source_revision', '未记录')}`。",
        f"- AIC 配置：`{metadata.get('simulator_config', '未记录')}`。",
        f"- CUDA Graph：`{metadata.get('real_cuda_graph_config', '未记录')}`。",
        f"- 内核：`{metadata.get('kernel_execution_path', '未记录')}`。",
        f"- AIC/真机 attention 对齐：`{metadata.get('attention_kernel_alignment', '未记录')}`。",
        "- 明细 CSV：`single_step_comparison.csv`；AIC 每算子预测分解在每个模型的 `aic_prediction.json`。",
        "",
    ]
    markdown_path = analysis / "single_step_comparison.md"
    markdown_path.write_text("\n".join(lines), encoding="utf-8")
    write_docx(analysis / "single_step_comparison.docx", rows, metadata)
    print(markdown_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
