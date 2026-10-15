#!/usr/bin/env python3
"""Render the small Markdown reports in this directory as self-contained DOCX files.

This intentionally handles the limited Markdown used by experiment reports:
headings, paragraphs, bullets, and pipe tables.  It avoids a Pandoc dependency
so the report can be regenerated in the hbf-cxl environment.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.image.image import Image
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def clean_inline(text: str) -> str:
    """Keep report text readable in Word without trying to be a full Markdown parser."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("`", "").replace("**", "").replace("__", "")


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def table_cells(line: str) -> list[str]:
    return [clean_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]


def set_simsun_rfonts(element) -> None:
    rpr = element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), "SimSun")


def set_simsun_run(run) -> None:
    run.font.name = "SimSun"
    set_simsun_rfonts(run._element)


def set_font(document: Document) -> None:
    """Explicitly set SimSun for every style used by generated reports."""
    for style in document.styles:
        try:
            style.font.name = "SimSun"
            set_simsun_rfonts(style._element)
        except AttributeError:
            continue
    document.styles["Normal"].font.size = Pt(9)


def apply_simsun(document: Document) -> None:
    """Cover direct formatting in headings, table cells, headers, and footers."""

    def paragraphs(container):
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)

    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))
    for container in containers:
        for paragraph in paragraphs(container):
            for run in paragraph.runs:
                set_simsun_run(run)


def add_table(document: Document, lines: list[str]) -> None:
    rows = [table_cells(line) for line in lines if not is_table_separator(line)]
    if not rows:
        return
    columns = len(rows[0])
    if any(len(row) != columns for row in rows):
        raise ValueError("Markdown table has inconsistent column counts")
    table = document.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, rows[0]):
        cell.text = value
    for values in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value


def render(markdown: Path, output: Path) -> None:
    lines = markdown.read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    set_font(document)
    document.core_properties.title = clean_inline(next((line[2:] for line in lines if line.startswith("# ")), markdown.stem))

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        if stripped.startswith("#"):
            match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
            if match:
                level = min(len(match.group(1)) - 1, 4)
                document.add_heading(clean_inline(match.group(2)), level=level)
                index += 1
                continue
        image = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if image:
            alt_text, target = image.groups()
            image_path = (markdown.parent / target).resolve()
            if image_path.is_file():
                if alt_text:
                    document.add_paragraph(clean_inline(alt_text))
                # Use the actual landscape page width, while leaving room
                # for the caption vertically; do not shrink to portrait size.
                picture = Image.from_file(str(image_path))
                available_width = section.page_width - section.left_margin - section.right_margin
                available_height = section.page_height - section.top_margin - section.bottom_margin - Inches(.4)
                scale = min(available_width / picture.width, available_height / picture.height)
                document.add_picture(str(image_path), width=int(picture.width * scale))
            else:
                document.add_paragraph(clean_inline(f"Figure unavailable: {target}"))
            index += 1
            continue
        if stripped.startswith("- "):
            document.add_paragraph(clean_inline(stripped[2:]), style="List Bullet")
        else:
            document.add_paragraph(clean_inline(stripped))
        index += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    apply_simsun(document)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    render(args.markdown.resolve(), args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
