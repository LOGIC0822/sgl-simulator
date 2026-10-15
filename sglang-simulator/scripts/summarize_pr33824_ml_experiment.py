#!/usr/bin/env python3
"""Build the final native-SGLang ML-predictor experiment handoff.

This combines three independently trained ``MLTimePredictor`` bundles with
the already completed five-repeat A100 serving baselines.  It rejects a result
root unless every selected model used predictor ``ml`` and every calibration
summary explicitly says that AIConfigurator was not used.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
from collections import defaultdict
from pathlib import Path
from statistics import mean, stdev
from typing import Any, Iterable


CORE_METRICS = (
    ("mean_ttft_ms", "TTFT (ms)"),
    ("mean_tpot_ms", "TPOT (ms)"),
    ("mean_e2e_latency_ms", "E2E (ms)"),
    ("total_throughput", "Total throughput (token/s)"),
)

COMPARISON_METHODS = (
    ("aic_degraded", "AIC：退化后端"),
    ("aic_compatible", "AIC：SM80 兼容 kernel"),
    ("ml_compatible", "ML：SM80 兼容 kernel"),
)
COMPARISON_PHASES = ("prefill", "decode")


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Cannot read JSON: {path}") from exc
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def load_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise ValueError(f"Missing CSV: {path}")
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def as_float(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def as_int(value: Any) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def fmt(value: Any, digits: int = 2) -> str:
    number = as_float(value)
    return "—" if number is None else f"{number:.{digits}f}"


def read_key_values(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.is_file():
        return values
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        key, sep, value = line.partition("=")
        if sep:
            values[key] = value
    return values


def write_csv(path: Path, rows: list[dict[str, Any]], fields: Iterable[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fields), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def aggregate_metric(rows: list[dict[str, str]], metric: str) -> dict[str, Any]:
    real = [as_float(row.get(f"real_{metric}")) for row in rows]
    predicted = [as_float(row.get(f"simulated_{metric}")) for row in rows]
    signed = [as_float(row.get(f"signed_error_pct_{metric}")) for row in rows]
    absolute = [as_float(row.get(f"absolute_error_pct_{metric}")) for row in rows]
    real_values = [value for value in real if value is not None]
    predicted_values = [value for value in predicted if value is not None]
    signed_values = [value for value in signed if value is not None]
    absolute_values = [value for value in absolute if value is not None]
    return {
        "paired_runs": len(absolute_values),
        "real_mean": mean(real_values) if real_values else None,
        "predicted_mean": mean(predicted_values) if predicted_values else None,
        "signed_error_mean_pct": mean(signed_values) if signed_values else None,
        "mape_pct": mean(absolute_values) if absolute_values else None,
        "absolute_error_min_pct": min(absolute_values) if absolute_values else None,
        "absolute_error_max_pct": max(absolute_values) if absolute_values else None,
        "absolute_error_std_pct": stdev(absolute_values) if len(absolute_values) > 1 else 0.0 if absolute_values else None,
    }


def calibration_rows(calibration_root: Path, models: list[str]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for model in models:
        summary_path = calibration_root / "models" / model / "analysis" / "training_summary.json"
        summary = load_json(summary_path)
        if summary.get("aiconfigurator_used") is not False:
            raise ValueError(f"Calibration summary does not prove AIC exclusion: {summary_path}")
        holdout = summary.get("holdout_metrics")
        if not isinstance(holdout, dict):
            raise ValueError(f"Missing holdout metrics: {summary_path}")
        for stage in ("all", "prefill", "decode"):
            metrics = holdout.get(stage)
            if not isinstance(metrics, dict) or as_int(metrics.get("n")) <= 0:
                raise ValueError(f"Missing {stage} holdout data: {summary_path}")
            result.append(
                {
                    "model": model,
                    "stage": stage,
                    "samples": as_int(metrics.get("n")),
                    "mape_pct": as_float(metrics.get("mape_pct")),
                    "median_ape_pct": as_float(metrics.get("median_ape_pct")),
                    "p90_ape_pct": as_float(metrics.get("p90_ape_pct")),
                    "mae_ms": as_float(metrics.get("mae_ms")),
                    "rmse_ms": as_float(metrics.get("rmse_ms")),
                    "mean_target_ms": as_float(metrics.get("mean_target_ms")),
                    "mean_prediction_ms": as_float(metrics.get("mean_prediction_ms")),
                    "model_file": summary.get("model_file", ""),
                    "simulator_config": summary.get("simulator_config", ""),
                }
            )
    return result


def calibration_telemetry_summary(calibration_root: Path) -> dict[str, Any]:
    """Summarize optional live GPU samples without inventing a clean-run claim.

    The collection launcher records the raw CSV when a long calibration is
    already in progress.  It may begin after an earlier model has finished, so
    consumers must treat it as evidence for its stated time window only.
    """
    path = calibration_root / "hardware" / "gpu_telemetry_live.csv"
    if not path.is_file():
        return {"available": False, "path": str(path)}
    rows = load_csv(path)
    if not rows:
        return {"available": False, "path": str(path)}

    utilization = [as_float(row.get("gpu_util_pct")) for row in rows]
    free_memory = [as_float(row.get("memory_free_mib")) for row in rows]
    app_counts = [
        len([item for item in row.get("compute_apps", "").split(";") if item.strip()])
        for row in rows
    ]
    return {
        "available": True,
        "path": str(path),
        "samples": len(rows),
        "first_timestamp": rows[0].get("timestamp", ""),
        "last_timestamp": rows[-1].get("timestamp", ""),
        "mean_gpu_util_pct": mean(value for value in utilization if value is not None)
        if any(value is not None for value in utilization)
        else None,
        "min_free_memory_mib": min(value for value in free_memory if value is not None)
        if any(value is not None for value in free_memory)
        else None,
        "samples_with_multiple_compute_processes": sum(count >= 2 for count in app_counts),
        "maximum_compute_processes": max(app_counts, default=0),
    }


def stage_row_count(stage_root: Path) -> int:
    path = stage_root / "result.jsonl"
    if not path.is_file():
        return 0
    return sum(1 for line in path.read_text(encoding="utf-8", errors="replace").splitlines() if line.strip())


def stage_expected_rows(stage_root: Path) -> int:
    try:
        value = int((stage_root / "expected_result_rows.txt").read_text(encoding="utf-8").strip())
    except (OSError, ValueError):
        return stage_row_count(stage_root)
    return value if value > 0 else stage_row_count(stage_root)


def calibration_grid_rows(
    calibration_root: Path,
    real_reference: Path,
    models: list[str],
) -> list[dict[str, Any]]:
    """Collect per-model shape coverage and cache settings for the report."""
    rows: list[dict[str, Any]] = []
    for model in models:
        model_root = calibration_root / "models" / model
        calibration = model_root / "calibration"
        holdout = model_root / "holdout"
        calibration_settings = read_key_values(model_root / "measurement-settings.txt")
        real_settings = read_key_values(real_reference / "results" / model / "real" / "launch-settings.txt")
        manifest = calibration / "grid-manifest.txt"
        rows.append(
            {
                "model": model,
                "calibration_rows": stage_row_count(calibration),
                "calibration_expected_rows": stage_expected_rows(calibration),
                "holdout_rows": stage_row_count(holdout),
                "holdout_expected_rows": stage_expected_rows(holdout),
                "calibration_mem_fraction_static": calibration_settings.get("mem_fraction_static", ""),
                "real_mem_fraction_static": real_settings.get("mem_fraction_static", ""),
                "grid_policy": "trace-feasible grouped grid" if manifest.is_file() else "static Cartesian grid",
            }
        )
    return rows


def collect_pairs(experiment_root: Path, models: list[str]) -> list[dict[str, str]]:
    pairs: list[dict[str, str]] = []
    for model in models:
        model_root = experiment_root / "per-model" / model
        if not (model_root / "MATRIX_COMPLETE").is_file():
            raise ValueError(f"ML matrix did not complete: {model_root}")
        metadata = read_key_values(model_root / "run_metadata.txt")
        if metadata.get("predictor_name") != "ml":
            raise ValueError(f"Result did not use ML predictor: {model_root}")
        config = load_json(model_root / "simulator-config.json")
        if (config.get("predictor") or {}).get("name") != "ml":
            raise ValueError(f"ML config not recorded at: {model_root}")
        rows = load_csv(model_root / "analysis" / "paired_repeat_metrics.csv")
        model_rows = [row for row in rows if row.get("model") == model and row.get("status") == "ok"]
        if len(model_rows) != 15:
            raise ValueError(f"Expected 15 complete paired rows for {model}, found {len(model_rows)}")
        for row in model_rows:
            row = dict(row)
            row["predictor"] = "ml"
            row["bundle_path"] = str((config.get("predictor") or {}).get("database_path", ""))
            pairs.append(row)
    return pairs


def make_summaries(pairs: list[dict[str, str]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    by_model: dict[str, list[dict[str, str]]] = defaultdict(list)
    by_case: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in pairs:
        by_model[row["model"]].append(row)
        by_case[(row["model"], row["case"])].append(row)

    model_rows: list[dict[str, Any]] = []
    for model, rows in sorted(by_model.items()):
        for metric, label in CORE_METRICS:
            model_rows.append({"model": model, "metric": metric, "metric_label": label, **aggregate_metric(rows, metric)})
    case_rows: list[dict[str, Any]] = []
    for (model, case), rows in sorted(by_case.items()):
        common = {
            "model": model,
            "case": case,
            "paired_runs": len(rows),
            "num_prompts": as_int(rows[0].get("num_prompts")),
            "input_tokens_target": as_int(rows[0].get("input_tokens_target")),
            "output_tokens_target": as_int(rows[0].get("output_tokens_target")),
            "request_rate": rows[0].get("request_rate", ""),
        }
        for metric, label in CORE_METRICS:
            case_rows.append({**common, "metric": metric, "metric_label": label, **aggregate_metric(rows, metric)})
    return model_rows, case_rows


def plot_figures(analysis: Path, model_rows: list[dict[str, Any]], operator_rows: list[dict[str, Any]]) -> list[Path]:
    try:
        import matplotlib.pyplot as plt
    except ImportError:
        return []
    figures = analysis / "figures"
    figures.mkdir(parents=True, exist_ok=True)
    models = sorted({row["model"] for row in model_rows})
    metrics = ["mean_ttft_ms", "mean_tpot_ms", "mean_e2e_latency_ms", "total_throughput"]
    labels = ["TTFT", "TPOT", "E2E", "Throughput"]
    width = 0.22
    fig, axis = plt.subplots(figsize=(10, 5.2))
    for index, model in enumerate(models):
        values = [next(row["mape_pct"] for row in model_rows if row["model"] == model and row["metric"] == metric) for metric in metrics]
        positions = [value + (index - (len(models) - 1) / 2) * width for value in range(len(metrics))]
        axis.bar(positions, values, width=width, label=model)
    axis.set_xticks(range(len(metrics)), labels)
    axis.set_ylabel("MAPE (%)")
    axis.set_title("End-to-end ML predictor error")
    axis.legend(fontsize=8)
    axis.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    e2e_path = figures / "e2e_mape_by_model.png"
    fig.savefig(e2e_path, dpi=180)
    plt.close(fig)

    stages = ["prefill", "decode"]
    fig, axis = plt.subplots(figsize=(8.5, 4.8))
    for index, stage in enumerate(stages):
        values = [next(row["mape_pct"] for row in operator_rows if row["model"] == model and row["stage"] == stage) for model in models]
        positions = [value + (index - 0.5) * width for value in range(len(models))]
        axis.bar(positions, values, width=width, label=stage)
    axis.set_xticks(range(len(models)), models, rotation=10)
    axis.set_ylabel("Held-out operator MAPE (%)")
    axis.set_title("Static-batch holdout error")
    axis.legend()
    axis.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    operator_path = figures / "operator_holdout_mape.png"
    fig.savefig(operator_path, dpi=180)
    plt.close(fig)
    return [e2e_path, operator_path]


def load_method_comparison(
    comparison_root: Path,
    models: list[str],
) -> dict[str, Any]:
    """Load the checked three-method comparison produced by the plotting script."""
    comparison_root = comparison_root.resolve()
    rows = load_csv(comparison_root / "comparison_data.csv")
    expected = {
        ("single_step", model, phase, method)
        for model in models
        for phase in COMPARISON_PHASES
        for method, _ in COMPARISON_METHODS
    }
    expected.update(
        ("end_to_end", model, "mean_e2e_latency_ms", method)
        for model in models
        for method, _ in COMPARISON_METHODS
    )
    indexed: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for row in rows:
        key = (
            row.get("comparison", ""),
            row.get("model", ""),
            row.get("phase", ""),
            row.get("method", ""),
        )
        if key not in expected:
            continue
        if key in indexed:
            raise ValueError(f"Duplicate comparison result {key}: {comparison_root}")
        error = as_float(row.get("absolute_error_pct"))
        if error is None or error < 0:
            raise ValueError(f"Invalid comparison MAPE for {key}: {comparison_root}")
        indexed[key] = {**row, "absolute_error_pct": error}
    if set(indexed) != expected:
        missing = sorted(expected - set(indexed))
        raise ValueError(f"Incomplete method comparison in {comparison_root}; missing={missing}")

    figures = [
        comparison_root / "single_step_mape.png",
        comparison_root / "end_to_end_e2e_mape.png",
    ]
    for figure in figures:
        if not figure.is_file():
            raise ValueError(f"Missing comparison figure: {figure}")
    return {"root": comparison_root, "rows": indexed, "figures": figures}


def comparison_table(
    comparison: dict[str, Any],
    models: list[str],
    scope: str,
    phase: str,
) -> list[str]:
    """Render one compact MAPE table from the already validated comparison rows."""
    rows = comparison["rows"]
    lines = [
        "| 模型 | AIC：退化后端 | AIC：SM80 兼容 kernel | ML：SM80 兼容 kernel |",
        "| --- | ---: | ---: | ---: |",
    ]
    for model in models:
        values = [
            rows[(scope, model, phase, method)]["absolute_error_pct"]
            for method, _ in COMPARISON_METHODS
        ]
        lines.append(f"| {model} | " + " | ".join(f"{fmt(value)}%" for value in values) + " |")
    return lines


def comparison_image_markdown(
    comparison: dict[str, Any],
    documentation_dir: Path | None,
) -> list[str]:
    """Return Markdown image references only when the report lives under docs/."""
    if documentation_dir is None:
        return []
    single_step, e2e = comparison["figures"]
    return [
        "",
        "![三种预测方式的固定单步误差对比]("
        + os.path.relpath(single_step, documentation_dir)
        + ")",
        "",
        "![三种预测方式的端到端误差对比]("
        + os.path.relpath(e2e, documentation_dir)
        + ")",
    ]


def markdown_report(
    calibration_root: Path,
    real_reference: Path,
    pairs: list[dict[str, str]],
    model_rows: list[dict[str, Any]],
    case_rows: list[dict[str, Any]],
    operator_rows: list[dict[str, Any]],
    telemetry: dict[str, Any],
    grid_rows: list[dict[str, Any]],
    comparison: dict[str, Any] | None = None,
    documentation_dir: Path | None = None,
) -> str:
    by_model_metric = {(row["model"], row["metric"]): row for row in model_rows}
    operator = {(row["model"], row["stage"]): row for row in operator_rows}
    models = sorted({row["model"] for row in model_rows})
    e2e_values = [row["mape_pct"] for row in model_rows if row["metric"] == "mean_e2e_latency_ms"]
    e2e_by_model = {
        row["model"]: row["mape_pct"]
        for row in model_rows
        if row["metric"] == "mean_e2e_latency_ms"
    }
    prediction_sample_counts = sorted(
        {
            as_int(row.get("prediction_samples"))
            for row in pairs
            if as_int(row.get("prediction_samples")) is not None
        }
    )
    if len(prediction_sample_counts) == 1:
        prediction_repeat_note = (
            f"本次每个工作负载的 ML 模拟均完成 {prediction_sample_counts[0]} 次，"
            "汇总时取其均值；真机侧也保留全部 5 次。"
        )
        value_description = (
            f"表中值为同一工作负载 5 次真机结果和 {prediction_sample_counts[0]} 次 "
            "ML 模拟结果的均值"
        )
    else:
        prediction_repeat_note = "ML 模拟按实际完成的重复次数汇总；真机侧保留全部 5 次。"
        value_description = "表中值为同一工作负载真机和 ML 模拟重复结果的均值"
    report = [
        "# A100 PCIe：原生 SGLang ML 预测器实验",
        "",
        "## 结果摘要",
        "",
        f"- 已完成 {len(models)} 个 dense 模型、3 组工作负载和 5 次真机基线，共 {len(pairs)} 组真实值/ML 预测值配对。",
        "- 本轮预测器是每个 checkpoint 单独训练的 sklearn ExtraTrees 模型；训练目标来自本机 A100 上原生 SGLang `one_batch` 的 CUDA 同步时延。训练、预测和汇总均未调用 AIConfigurator。",
        f"- 三个模型的 E2E MAPE 平均为 {fmt(mean(e2e_values))}%；具体真实值、预测值和误差在下表逐项列出。",
        "- 本轮端到端 E2E MAPE 分别为 "
        + "；".join(f"{model} {fmt(e2e_by_model[model])}%" for model in models)
        + "。",
        "- 静态单批 holdout 和端到端矩阵是两类测量：前者使用留出的 batch/sequence 形状，后者比较固定请求 trace 的服务结果。",
        "",
        "## 实验边界",
        "",
        f"- 真机基线：`{real_reference}`，A100 PCIe、完整 checkpoint、原生 SGLang、Triton attention、piecewise CUDA graph 和 overlap scheduler；每个固定 trace 跑 5 次。",
        f"- ML 校准：`{calibration_root}`。3B/7B 使用静态 Cartesian 网格；32B 使用按真实 trace 的 prefill 上限和 decode 并发组织的可行分组网格。",
        "- 模拟：PR #33824 CPU 离线 scheduler，`MLTimePredictor` 仅预测每次 forward 的逻辑时长，不执行物理 CUDA forward。" + prediction_repeat_note,
        "",
        "| 模型 | 校准形状 | 留出形状 | 校准 mem_fraction | 真机基线 mem_fraction | 采样方式 |",
        "| --- | ---: | ---: | ---: | ---: | --- |",
    ]
    for row in grid_rows:
        report.append(
            f"| {row['model']} | {row['calibration_rows']} / {row['calibration_expected_rows']} | "
            f"{row['holdout_rows']} / {row['holdout_expected_rows']} | "
            f"{row['calibration_mem_fraction_static'] or '—'} | "
            f"{row['real_mem_fraction_static'] or '—'} | {row['grid_policy']} |"
        )
    report.extend(
        [
            "",
            "- 32B 的最初 32×1024 静态 prefill 在 80GB 上触发 OOM，且该形状不在固定服务 trace 中：trace 的单轮 prefill 总扩展 token 不超过 2048。失败日志已保留，最终 32B 结果改为同样 `mem_fraction_static=0.96` 下的 trace-feasible 网格。",
            "",
            "## 静态批次 holdout",
            "",
            "| 模型 | 阶段 | 样本数 | MAPE | 中位 APE | P90 APE | MAE (ms) |",
            "| --- | --- | ---: | ---: | ---: | ---: | ---: |",
        ]
    )
    for model in models:
        for stage in ("prefill", "decode"):
            row = operator[(model, stage)]
            report.append(
                f"| {model} | {stage} | {row['samples']} | {fmt(row['mape_pct'])}% | "
                f"{fmt(row['median_ape_pct'])}% | {fmt(row['p90_ape_pct'])}% | {fmt(row['mae_ms'])} |"
            )
    if comparison is not None:
        comparison_root = comparison["root"]
        report.extend(
            [
                "",
                "## 与 AIC 预测的并列对照",
                "",
                "该节使用已生成的三方法比较表。单步均为 batch=1、1024-token prefill 和首个 decode token；端到端均为 3 个负载、每负载 5 次真实/预测配对。",
                "",
                "### 单步 Prefill MAPE",
                "",
                *comparison_table(comparison, models, "single_step", "prefill"),
                "",
                "### 单步 Decode MAPE",
                "",
                *comparison_table(comparison, models, "single_step", "decode"),
                "",
                "### 端到端 E2E MAPE",
                "",
                *comparison_table(comparison, models, "end_to_end", "mean_e2e_latency_ms"),
                "",
                "- AIC 退化后端和 AIC SM80 兼容 kernel 分别与各自批次的真实基线配对；两批运行不是只改变一个 kernel 的受控对照。",
                "- “SM80 兼容 kernel”描述真机负载使用的兼容包；AIC 查询仍使用 Silicon 静态算子表，不直接执行 CUDA kernel。",
                "- ML 的固定单步点位使用兼容 kernel 的真机均值。input=1024 位于 ML 校准网格内，因此这两个单点不是独立 holdout。",
                "- ML 的端到端结果复用兼容 kernel 矩阵的真机基线。",
                f"- 图表、逐项数据和来源：`{comparison_root}`。",
                *comparison_image_markdown(comparison, documentation_dir),
            ]
        )
    if telemetry.get("available"):
        report.extend(
            [
                "",
                "## 采集期间的 GPU 状态",
                "",
                f"- 动态遥测从 {telemetry['first_timestamp']} 到 {telemetry['last_timestamp']}，共 {telemetry['samples']} 个两秒采样点；平均 GPU 利用率为 {fmt(telemetry['mean_gpu_util_pct'])}%，最小空闲显存为 {fmt(telemetry['min_free_memory_mib'])} MiB。",
                f"- 其中 {telemetry['samples_with_multiple_compute_processes']} 个采样点同时发现至少两个 GPU compute 进程（最多 {telemetry['maximum_compute_processes']} 个）。遥测在 3B 采集完成后才开始，因此它只能说明后续时间窗口的资源竞争，不能反证 3B 采集是否干净。",
            ]
        )
    report.extend(
        [
            "",
            "## 端到端模型汇总",
            "",
            "| 模型 | TTFT MAPE | TPOT MAPE | E2E MAPE | 总吞吐 MAPE |",
            "| --- | ---: | ---: | ---: | ---: |",
        ]
    )
    for model in models:
        values = [by_model_metric[(model, metric)]["mape_pct"] for metric, _ in CORE_METRICS]
        report.append(f"| {model} | " + " | ".join(f"{fmt(value)}%" for value in values) + " |")
    report.extend(
        [
            "",
            "## 端到端真实值与预测值",
            "",
            value_description + "；signed error 为 `(预测−真实)/真实×100%`。",
            "",
            "| 模型 | 工作负载 | 指标 | 真实均值 | ML 预测 | signed error | MAPE |",
            "| --- | --- | --- | ---: | ---: | ---: | ---: |",
        ]
    )
    for row in case_rows:
        report.append(
            f"| {row['model']} | {row['case']} | {row['metric_label']} | {fmt(row['real_mean'])} | "
            f"{fmt(row['predicted_mean'])} | {fmt(row['signed_error_mean_pct'])}% | {fmt(row['mape_pct'])}% |"
        )
    report.extend(
        [
            "",
            "## 说明",
            "",
            "- 每个模型的 ML bundle 只使用本模型的静态真机采样；没有用这 45 条端到端真实基线拟合或调参。",
            "- 模型输入是 `MLTimePredictor` 的 18 个调度批次长度特征。它不包含模型结构字段，因此不能把一个模型的 bundle 直接套到另一个 checkpoint。",
            "- 结果文件中的 `per-model/<model>/analysis/paired_repeat_metrics.csv` 保留逐次真实值与预测值；`analysis/operator_holdout_summary.csv` 保留单步 holdout 误差。",
        ]
    )
    return "\n".join(report) + "\n"


def set_simsun(run: Any) -> None:
    from docx.oxml.ns import qn

    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run._element.rPr.rFonts.set(qn("w:ascii"), "SimSun")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "SimSun")


def write_docx(path: Path, markdown: str, figures: list[Path]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:ascii"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "SimSun")
    document.styles["Normal"].font.size = Pt(10.5)
    for section in document.sections:
        for paragraph in (section.header.paragraphs[0], section.footer.paragraphs[0]):
            paragraph.text = "SGLang Simulator ML predictor experiment"
            for run in paragraph.runs:
                set_simsun(run)
                run.font.size = Pt(8)
    title = document.add_heading("A100 PCIe：原生 SGLang ML 预测器实验", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        set_simsun(run)

    rendered_figures: set[Path] = set()
    for line in markdown.splitlines()[2:]:
        if line.startswith("### "):
            heading = document.add_heading(line[4:], level=2)
            for run in heading.runs:
                set_simsun(run)
        elif line.startswith("## "):
            heading = document.add_heading(line[3:], level=1)
            for run in heading.runs:
                set_simsun(run)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(line[2:], style="List Bullet")
            for run in paragraph.runs:
                set_simsun(run)
        elif line.startswith("|") and not line.startswith("| ---"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if not hasattr(write_docx, "_table") or getattr(write_docx, "_width", 0) != len(cells):
                write_docx._table = document.add_table(rows=0, cols=len(cells))  # type: ignore[attr-defined]
                write_docx._width = len(cells)  # type: ignore[attr-defined]
            table = write_docx._table  # type: ignore[attr-defined]
            row = table.add_row().cells
            for cell, text in zip(row, cells, strict=True):
                cell.text = text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_simsun(run)
        elif line.startswith("!["):
            alt_text, _, target = line[2:].partition("](")
            figure = next((item for item in figures if item.name in target), None)
            if figure is not None:
                paragraph = document.add_paragraph(alt_text)
                for run in paragraph.runs:
                    set_simsun(run)
                document.add_picture(str(figure), width=Inches(6.2))
                rendered_figures.add(figure)
        elif line and not line.startswith("#"):
            paragraph = document.add_paragraph(line)
            for run in paragraph.runs:
                set_simsun(run)
        elif not line:
            if hasattr(write_docx, "_table"):
                delattr(write_docx, "_table")
                delattr(write_docx, "_width")
    for figure in figures:
        if figure in rendered_figures:
            continue
        document.add_picture(str(figure), width=Inches(6.2))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--experiment-root", type=Path, required=True)
    parser.add_argument("--calibration-root", type=Path, required=True)
    parser.add_argument("--real-reference-root", type=Path, required=True)
    parser.add_argument("--models", required=True)
    parser.add_argument(
        "--documentation-dir",
        type=Path,
        help="Optional docs directory receiving a Markdown/DOCX copy of the final report.",
    )
    parser.add_argument(
        "--comparison-root",
        type=Path,
        help="Optional output directory from plot_pr33824_predictor_method_comparison.py.",
    )
    args = parser.parse_args()
    models = [value.strip() for value in args.models.split(",") if value.strip()]
    if not models:
        raise SystemExit("--models cannot be empty")
    root = args.experiment_root.resolve()
    analysis = root / "analysis"
    resolved_calibration_root = args.calibration_root.resolve()
    operator_rows = calibration_rows(resolved_calibration_root, models)
    telemetry = calibration_telemetry_summary(resolved_calibration_root)
    grid_rows = calibration_grid_rows(
        resolved_calibration_root,
        args.real_reference_root.resolve(),
        models,
    )
    comparison = (
        load_method_comparison(args.comparison_root, models)
        if args.comparison_root is not None
        else None
    )
    pairs = collect_pairs(root, models)
    model_rows, case_rows = make_summaries(pairs)
    expected_pairs = len(models) * 3 * 5
    if len(pairs) != expected_pairs:
        raise SystemExit(f"Expected {expected_pairs} pairs, found {len(pairs)}")
    pair_fields = list(pairs[0])
    write_csv(analysis / "paired_repeat_metrics.csv", pairs, pair_fields)
    write_csv(analysis / "model_metric_summary.csv", model_rows, model_rows[0].keys())
    write_csv(analysis / "case_metric_summary.csv", case_rows, case_rows[0].keys())
    write_csv(analysis / "operator_holdout_summary.csv", operator_rows, operator_rows[0].keys())
    figures = plot_figures(analysis, model_rows, operator_rows)
    if comparison is not None:
        figures.extend(comparison["figures"])
    report = markdown_report(
        resolved_calibration_root,
        args.real_reference_root.resolve(),
        pairs,
        model_rows,
        case_rows,
        operator_rows,
        telemetry,
        grid_rows,
        comparison,
    )
    (analysis / "final_report.md").write_text(report, encoding="utf-8")
    write_docx(analysis / "final_report.docx", report, figures)
    documentation_report: dict[str, str] | None = None
    if args.documentation_dir is not None:
        documentation_dir = args.documentation_dir.resolve()
        experiment_id = root.name.removeprefix("pr33824_ml_dense3_matrix_")
        report_stem = f"a100_pcie_ml_predictor_dense3_{experiment_id}"
        documentation_markdown = documentation_dir / f"{report_stem}.md"
        documentation_docx = documentation_dir / f"{report_stem}.docx"
        documentation_markdown.parent.mkdir(parents=True, exist_ok=True)
        documentation_report_text = markdown_report(
            resolved_calibration_root,
            args.real_reference_root.resolve(),
            pairs,
            model_rows,
            case_rows,
            operator_rows,
            telemetry,
            grid_rows,
            comparison,
            documentation_dir,
        )
        documentation_markdown.write_text(documentation_report_text, encoding="utf-8")
        write_docx(documentation_docx, documentation_report_text, figures)
        documentation_report = {
            "markdown": str(documentation_markdown),
            "docx": str(documentation_docx),
        }
    summary = {
        "status": "complete",
        "predictor": "ml",
        "aiconfigurator_used": False,
        "models": models,
        "expected_paired_runs": expected_pairs,
        "completed_paired_runs": len(pairs),
        "calibration_root": str(args.calibration_root.resolve()),
        "real_reference_root": str(args.real_reference_root.resolve()),
        "calibration_gpu_telemetry": telemetry,
        "calibration_grid": grid_rows,
        "comparison_root": str(comparison["root"]) if comparison is not None else None,
        "report": str((analysis / "final_report.md").resolve()),
        "documentation_report": documentation_report,
    }
    (analysis / "summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
