#!/usr/bin/env python3
"""Locate AIC-versus-SGLang error sources in an existing CUDA Graph trace.

This is a post-processing tool.  It does not launch a model or modify an AIC
database.  It reads the graph-enabled ``one_batch`` Chrome traces, re-queries
the same calibrated AIC database for the attention internals, and keeps only
kernel families whose origin can be identified from the trace name and the
executed SGLang source path.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from aiconfigurator.sdk.operations.attention import (
    ContextAttention,
    GenerationAttention,
)
from docx import Document
from docx.shared import Pt
from sglang.srt.configs.model_config import ModelConfig

from sglang_simulator.simulation.sglang.utils import resolve_model_info
from sglang_simulator.simulation.types import SchedulerConfig
from sglang_simulator.spec.accelerator.base import AcceleratorInfo
from sglang_simulator.spec.data_type import DataType
from sglang_simulator.time_predictor.aiconfigurator import AIConfiguratorTimePredictor


MODELS = ("Qwen2.5-3B-Instruct", "Qwen2.5-7B-Instruct", "Qwen3-32B")
PHASES = ("prefill", "decode")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_key_value_file(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if "=" in line:
            key, value = line.split("=", 1)
            result[key] = value
    return result


def duration_us(event: dict[str, Any]) -> float:
    value = event.get("dur", 0.0)
    return float(value) if isinstance(value, (int, float)) else 0.0


def event_grid(event: dict[str, Any]) -> tuple[int, ...]:
    grid = event.get("args", {}).get("grid", [])
    return tuple(int(value) for value in grid if isinstance(value, (int, float)))


def is_attention_core(name: str, phase: str) -> bool:
    lowered = name.lower()
    if phase == "prefill":
        return lowered == "_fwd_kernel"
    return (
        lowered == "_fwd_kernel"
        or lowered.startswith("_fwd_grouped_kernel_stage")
        or lowered.startswith("_fwd_kernel_stage")
    )


def raw_kernel_family(
    event: dict[str, Any], phase: str, silu_grids: set[tuple[int, ...]]
) -> str:
    """Return a conservative family for one trace event.

    The names selected here have a direct code-path interpretation:

    * the FP32 ``pow``/``mean``/``rsqrt`` chain is native RMSNorm;
    * a SiLU grid plus a same-grid BF16 multiply is native ``SiluAndMul``;
    * the remaining BF16 multiplication grids match Q/K RoPE tensor shapes.

    Ambiguous residual adds and generic copies intentionally remain in
    ``residual_cast_misc``.  The resulting norm and attention-auxiliary totals
    are therefore conservative lower bounds, not an attempt to guess every
    fused high-level operation from graph replay order.
    """

    category = str(event.get("cat", ""))
    name = str(event.get("name", ""))
    lowered = name.lower()

    if category == "gpu_memcpy":
        return "memcpy"
    if category == "gpu_memset":
        return "memset"
    if "gemm" in lowered or "gemv" in lowered or "cublaslt" in lowered:
        return "gemm"
    if is_attention_core(name, phase):
        return "attention_fmha_core"
    if "silu_kernel" in lowered:
        return "swiglu_silu"
    if "binaryfunctor<c10::bfloat16" in lowered and "mulf" in lowered:
        return (
            "swiglu_gate_mul"
            if event_grid(event) in silu_grids
            else "rope_shaped_bf16_mul"
        )
    if (
        "pow_tensor_scalar_kernel_impl<float" in lowered
        or "meanops<float" in lowered
        or "rsqrt_kernel_cuda" in lowered
        or "cudafunctoronself_add<float>" in lowered
        or (
            "binaryfunctor<float, float, float" in lowered and "mulf" in lowered
        )
        or "cudafunctor_add<float>" in lowered
    ):
        return "rmsnorm_fp32_core"
    if "direct_copy_kernel_cuda" in lowered and "lambda(float)" in lowered:
        return "rmsnorm_bf16_to_fp32"
    if "bfloat16_copy_kernel_cuda" in lowered:
        return "rmsnorm_fp32_to_bf16"
    if "fused_qknorm_warp" in lowered:
        return "qk_norm"
    if "catarraybatchedcopy" in lowered:
        return "rope_cat_copy"
    if "indexselectsmallindex<float" in lowered or "vectorized_gather_kernel" in lowered:
        return "rope_position_gather"
    if "store_kvcache" in lowered:
        return "kv_store"
    if any(
        part in lowered
        for part in (
            "create_flashinfer_kv_indices",
            "get_num_kv_splits",
            "write_req_to_token_pool",
            "compute_position_kernel",
            "clamp_position_kernel",
            "devicescankernel",
            "devicescaninitkernel",
            "gpu_index_kernel",
        )
    ):
        return "cache_index"
    return "residual_cast_misc"


def trace_profile(path: Path, phase: str) -> dict[str, Any]:
    with gzip.open(path, "rt", encoding="utf-8") as file:
        events = json.load(file).get("traceEvents", [])
    gpu_events = [
        event
        for event in events
        if event.get("cat") in {"kernel", "gpu_memcpy", "gpu_memset"}
        and duration_us(event) > 0
    ]
    silu_grids = {
        event_grid(event)
        for event in gpu_events
        if "silu_kernel" in str(event.get("name", "")).lower()
    }
    totals_us: Counter[str] = Counter()
    counts: Counter[str] = Counter()
    for event in gpu_events:
        family = raw_kernel_family(event, phase, silu_grids)
        totals_us[family] += duration_us(event)
        counts[family] += 1

    family_ms = {name: value / 1000.0 for name, value in sorted(totals_us.items())}
    total_ms = sum(family_ms.values())
    rmsnorm_ms = sum(
        family_ms.get(name, 0.0)
        for name in (
            "rmsnorm_fp32_core",
            "rmsnorm_bf16_to_fp32",
            "rmsnorm_fp32_to_bf16",
        )
    )
    swiglu_ms = family_ms.get("swiglu_silu", 0.0) + family_ms.get(
        "swiglu_gate_mul", 0.0
    )
    attention_aux_ms = sum(
        family_ms.get(name, 0.0)
        for name in (
            "rope_shaped_bf16_mul",
            "rope_cat_copy",
            "rope_position_gather",
            "qk_norm",
            "kv_store",
        )
    )
    return {
        "trace_gpu_sum_ms": total_ms,
        "families_ms": family_ms,
        "families_events": dict(counts),
        "gemm_ms": family_ms.get("gemm", 0.0),
        "attention_fmha_core_ms": family_ms.get("attention_fmha_core", 0.0),
        "attention_aux_identified_ms": attention_aux_ms,
        "attention_identified_total_ms": family_ms.get("attention_fmha_core", 0.0)
        + attention_aux_ms,
        "native_rmsnorm_identified_ms": rmsnorm_ms,
        "native_swiglu_ms": swiglu_ms,
    }


def error_percent(predicted: float, measured: float) -> float | None:
    return (predicted - measured) / measured * 100.0 if measured else None


def fmt(value: float | None, digits: int = 3) -> str:
    return "—" if value is None else f"{value:.{digits}f}"


def make_predictor(
    model_path: Path, simulator_config: dict[str, Any]
) -> AIConfiguratorTimePredictor:
    accelerator_config = simulator_config["platform"]["accelerator"]
    scheduler = simulator_config["scheduler"]
    accelerator = AcceleratorInfo.from_dict(
        {
            "name": str(accelerator_config["name"]),
            "vendor": str(accelerator_config.get("vendor", "NVIDIA")),
            "hbm_capacity_gb": int(accelerator_config["hbm_capacity_gb"]),
            "hbm_bandwidth_gb": int(accelerator_config["hbm_bandwidth_gb"]),
            "intra_node_bandwidth_gb": int(
                accelerator_config.get("intra_node_bandwidth_gb", 32)
            ),
            "inter_node_bandwidth_gb": int(
                accelerator_config.get("inter_node_bandwidth_gb", 25)
            ),
        }
    )
    scheduler_config = SchedulerConfig(
        data_type=DataType.BF16,
        kv_cache_data_type=DataType.BF16,
        tp_size=int(scheduler.get("tp_size", 1)),
        ep_size=int(scheduler.get("ep_size", 1)),
        dp_size=int(scheduler.get("dp_size", 1)),
        pp_size=int(scheduler.get("pp_size", 1)),
        backend_name=str(scheduler.get("backend_name", "sglang")),
        backend_version=str(scheduler.get("backend_version", "0.5.10")),
    )
    model_config = ModelConfig(model_path=str(model_path), trust_remote_code=False)
    model = resolve_model_info(model_config)
    return AIConfiguratorTimePredictor(
        model,
        hw=accelerator,
        config=scheduler_config,
        database_path=str(Path(simulator_config["predictor"]["database_path"])),
        database_mode=str(simulator_config["predictor"].get("database_mode", "SILICON")),
    )


def aic_attention_parts(
    predictor: AIConfiguratorTimePredictor,
    input_len: int,
) -> dict[str, Any]:
    """Expose the AIC attention table versus its analytic auxiliary terms."""

    database = predictor._session._database
    model = predictor._session._model
    context = next(
        operation
        for operation in model.context_ops
        if isinstance(operation, ContextAttention)
    )
    generation = next(
        operation
        for operation in model.generation_ops
        if isinstance(operation, GenerationAttention)
    )

    context_table_per_layer = context._query_context_attention_table(
        database,
        1,
        input_len,
        0,
        context._n,
        context._n_kv,
        context._kvcache_quant_mode,
        context._fmha_quant_mode,
        window_size=context._window_size,
        head_size=context._head_size,
    )
    context_table_ms = float(context_table_per_layer) * float(context._scale_factor)
    context_total_ms = float(
        context.query(
            database,
            x=input_len,
            batch_size=1,
            beam_width=1,
            s=input_len,
            prefix=0,
            seq_imbalance_correction_scale=1.0,
        )
    )

    generation_table_per_layer = generation._query_generation_attention_table(
        database,
        1,
        input_len + 1,
        generation._n,
        generation._n_kv,
        generation._kv_cache_dtype,
        window_size=generation._window_size,
        head_size=generation._head_size,
    )
    generation_table_ms = float(generation_table_per_layer) * float(
        generation._scale_factor
    )
    generation_total_ms = float(
        generation.query(
            database,
            x=1,
            batch_size=1,
            beam_width=1,
            s=input_len + 1,
            gen_seq_imbalance_correction_scale=1.0,
        )
    )

    return {
        "prefill": {
            "table_core_ms": context_table_ms,
            "analytic_aux_ms": context_total_ms - context_table_ms,
            "total_ms": context_total_ms,
            "use_qk_norm": bool(context._use_qk_norm),
        },
        "decode": {
            "table_core_ms": generation_table_ms,
            "analytic_aux_ms": generation_total_ms - generation_table_ms,
            "total_ms": generation_total_ms,
            "use_qk_norm": bool(generation._use_qk_norm),
        },
    }


def aic_breakdown(prediction: dict[str, Any], phase: str) -> dict[str, float]:
    breakdown = prediction["phases"][phase]["breakdown_ms"]
    return {
        "total_ms": float(prediction["phases"][phase]["predicted_ms"]),
        "gemm_ms": sum(
            float(value) for name, value in breakdown.items() if name.endswith("_gemm")
        ),
        "add_norm_ms": sum(
            float(value)
            for name, value in breakdown.items()
            if name.endswith("_add_norm_1") or name.endswith("_add_norm_2")
        ),
        "act_gate_ms": sum(
            float(value) for name, value in breakdown.items() if name.endswith("_act_gate")
        ),
    }


def add_docx_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results-root", type=Path, required=True)
    parser.add_argument("--baseline-root", type=Path, required=True)
    parser.add_argument("--checkpoint-root", type=Path, default=Path("/mnt/data/models"))
    parser.add_argument("--models", default=",".join(MODELS))
    args = parser.parse_args()

    root = args.results_root.resolve()
    baseline = args.baseline_root.resolve()
    metadata = read_key_value_file(root / "run_metadata.txt")
    baseline_metadata = read_key_value_file(baseline / "run_metadata.txt")
    input_len = int(metadata["input_len"])
    simulator_config = read_json(Path(baseline_metadata["simulator_config"]))
    models = tuple(filter(None, args.models.split(",")))
    analysis = root / "analysis"
    analysis.mkdir(parents=True, exist_ok=True)

    phase_rows: list[dict[str, Any]] = []
    component_rows: list[dict[str, Any]] = []
    raw_family_rows: list[dict[str, Any]] = []
    attention_rows: list[dict[str, Any]] = []
    profile_json: dict[str, Any] = {"metadata": metadata, "models": {}}

    for model_name in models:
        model_path = args.checkpoint_root / model_name
        if not model_path.is_dir():
            raise FileNotFoundError(f"Checkpoint not found: {model_path}")
        prediction = read_json(
            baseline / "models" / model_name / "aic_prediction.json"
        )
        attention_parts = aic_attention_parts(
            make_predictor(model_path, simulator_config), input_len
        )
        profile_json["models"][model_name] = {"phases": {}}

        for phase in PHASES:
            traces = sorted(
                (root / "models" / model_name / "traces").glob(
                    f"*_{phase}.trace.json.gz"
                )
            )
            if len(traces) != 1:
                raise RuntimeError(
                    f"Expected one {phase} trace for {model_name}, found {traces}"
                )
            trace = trace_profile(traces[0], phase)
            aic = aic_breakdown(prediction, phase)
            attention = attention_parts[phase]
            row = {
                "model": model_name,
                "phase": phase,
                "aic_total_ms": aic["total_ms"],
                "trace_gpu_sum_ms": trace["trace_gpu_sum_ms"],
                "total_error_pct": error_percent(
                    aic["total_ms"], trace["trace_gpu_sum_ms"]
                ),
                "aic_gemm_ms": aic["gemm_ms"],
                "trace_gemm_ms": trace["gemm_ms"],
                "gemm_error_pct": error_percent(aic["gemm_ms"], trace["gemm_ms"]),
                "aic_add_norm_ms": aic["add_norm_ms"],
                "trace_native_rmsnorm_identified_ms": trace[
                    "native_rmsnorm_identified_ms"
                ],
                "add_norm_error_pct": error_percent(
                    aic["add_norm_ms"], trace["native_rmsnorm_identified_ms"]
                ),
                "aic_act_gate_ms": aic["act_gate_ms"],
                "trace_native_swiglu_ms": trace["native_swiglu_ms"],
                "act_gate_error_pct": error_percent(
                    aic["act_gate_ms"], trace["native_swiglu_ms"]
                ),
                "aic_attention_table_core_ms": attention["table_core_ms"],
                "trace_attention_fmha_core_ms": trace["attention_fmha_core_ms"],
                "attention_core_error_pct": error_percent(
                    attention["table_core_ms"], trace["attention_fmha_core_ms"]
                ),
                "aic_attention_aux_ms": attention["analytic_aux_ms"],
                "trace_attention_aux_identified_ms": trace[
                    "attention_aux_identified_ms"
                ],
                "attention_aux_error_pct": error_percent(
                    attention["analytic_aux_ms"],
                    trace["attention_aux_identified_ms"],
                ),
                "aic_attention_total_ms": attention["total_ms"],
                "trace_attention_identified_total_ms": trace[
                    "attention_identified_total_ms"
                ],
                "attention_total_lower_bound_error_pct": error_percent(
                    attention["total_ms"], trace["attention_identified_total_ms"]
                ),
                "aic_use_qk_norm": attention["use_qk_norm"],
                "trace": str(traces[0]),
            }
            observed_gap_ms = row["trace_gpu_sum_ms"] - row["aic_total_ms"]
            gemm_gap_ms = row["trace_gemm_ms"] - row["aic_gemm_ms"]
            attention_gap_ms = (
                row["trace_attention_identified_total_ms"]
                - row["aic_attention_total_ms"]
            )
            rmsnorm_gap_ms = (
                row["trace_native_rmsnorm_identified_ms"] - row["aic_add_norm_ms"]
            )
            swiglu_gap_ms = row["trace_native_swiglu_ms"] - row["aic_act_gate_ms"]
            identified_gap_ms = (
                gemm_gap_ms + attention_gap_ms + rmsnorm_gap_ms + swiglu_gap_ms
            )
            row.update(
                {
                    "observed_gpu_gap_ms": observed_gap_ms,
                    "gemm_gap_ms": gemm_gap_ms,
                    "attention_gap_lower_bound_ms": attention_gap_ms,
                    "rmsnorm_gap_ms": rmsnorm_gap_ms,
                    "swiglu_gap_ms": swiglu_gap_ms,
                    "identified_gap_ms": identified_gap_ms,
                    "remaining_gpu_gap_ms": observed_gap_ms - identified_gap_ms,
                    "identified_gap_share_pct": (
                        identified_gap_ms / observed_gap_ms * 100.0
                        if observed_gap_ms
                        else None
                    ),
                }
            )
            phase_rows.append(row)
            profile_json["models"][model_name]["phases"][phase] = {
                "summary": row,
                "trace": trace,
                "aic_attention": attention,
            }

            components = (
                (
                    "gemm",
                    aic["gemm_ms"],
                    trace["gemm_ms"],
                    "AIC calibrated GEMM rows vs cuBLAS/cuBLASLt trace kernels",
                ),
                (
                    "attention_fmha_core",
                    attention["table_core_ms"],
                    trace["attention_fmha_core_ms"],
                    "AIC attention table core vs Triton FMHA kernel(s)",
                ),
                (
                    "attention_auxiliary",
                    attention["analytic_aux_ms"],
                    trace["attention_aux_identified_ms"],
                    "AIC analytic RoPE/QK-norm/KV term vs identified trace lower bound",
                ),
                (
                    "attention_identified_total",
                    attention["total_ms"],
                    trace["attention_identified_total_ms"],
                    "AIC attention operation vs FMHA plus identified auxiliary lower bound",
                ),
                (
                    "add_norm",
                    aic["add_norm_ms"],
                    trace["native_rmsnorm_identified_ms"],
                    "AIC ElementWise add+norm vs native Torch RMSNorm FP32/cast chain",
                ),
                (
                    "act_gate",
                    aic["act_gate_ms"],
                    trace["native_swiglu_ms"],
                    "AIC ElementWise gate vs native Torch SiLU plus same-grid BF16 multiply",
                ),
            )
            for component, aic_ms, trace_ms, basis in components:
                component_rows.append(
                    {
                        "model": model_name,
                        "phase": phase,
                        "component": component,
                        "aic_ms": aic_ms,
                        "trace_ms": trace_ms,
                        "error_pct": error_percent(aic_ms, trace_ms),
                        "basis": basis,
                    }
                )
            for family, milliseconds in trace["families_ms"].items():
                raw_family_rows.append(
                    {
                        "model": model_name,
                        "phase": phase,
                        "family": family,
                        "trace_ms": milliseconds,
                        "gpu_event_count": trace["families_events"].get(family, 0),
                        "trace_gpu_share_pct": milliseconds
                        / trace["trace_gpu_sum_ms"]
                        * 100.0,
                    }
                )
            attention_rows.append(
                {
                    "model": model_name,
                    "phase": phase,
                    "use_qk_norm": attention["use_qk_norm"],
                    "aic_table_core_ms": attention["table_core_ms"],
                    "aic_analytic_aux_ms": attention["analytic_aux_ms"],
                    "aic_attention_total_ms": attention["total_ms"],
                    "trace_fmha_core_ms": trace["attention_fmha_core_ms"],
                    "trace_aux_identified_lower_bound_ms": trace[
                        "attention_aux_identified_ms"
                    ],
                    "trace_attention_identified_lower_bound_ms": trace[
                        "attention_identified_total_ms"
                    ],
                }
            )

    write_csv(analysis / "aic_raw_kernel_phase_summary.csv", phase_rows)
    write_csv(analysis / "aic_raw_kernel_component_alignment.csv", component_rows)
    write_csv(analysis / "aic_attention_internal_breakdown.csv", attention_rows)
    write_csv(analysis / "raw_cuda_kernel_families.csv", raw_family_rows)
    (analysis / "aic_raw_kernel_root_cause.json").write_text(
        json.dumps(profile_json, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    aggregate_table = [
        [
            row["model"],
            row["phase"],
            fmt(row["aic_total_ms"]),
            fmt(row["trace_gpu_sum_ms"]),
            fmt(row["total_error_pct"], 2) + "%",
        ]
        for row in phase_rows
    ]
    elementwise_table = [
        [
            row["model"],
            row["phase"],
            fmt(row["aic_add_norm_ms"]),
            fmt(row["trace_native_rmsnorm_identified_ms"]),
            fmt(row["add_norm_error_pct"], 1) + "%",
            fmt(row["aic_act_gate_ms"]),
            fmt(row["trace_native_swiglu_ms"]),
            fmt(row["act_gate_error_pct"], 1) + "%",
        ]
        for row in phase_rows
    ]
    attention_table = [
        [
            row["model"],
            row["phase"],
            fmt(row["aic_attention_table_core_ms"]),
            fmt(row["trace_attention_fmha_core_ms"]),
            fmt(row["attention_core_error_pct"], 1) + "%",
            fmt(row["aic_attention_aux_ms"]),
            fmt(row["trace_attention_aux_identified_ms"]),
            fmt(row["attention_aux_error_pct"], 1) + "%",
            fmt(row["aic_attention_total_ms"]),
            fmt(row["trace_attention_identified_total_ms"]),
            fmt(row["attention_total_lower_bound_error_pct"], 1) + "%",
        ]
        for row in phase_rows
    ]
    gap_table = [
        [
            row["model"],
            row["phase"],
            fmt(row["observed_gpu_gap_ms"]),
            fmt(row["gemm_gap_ms"]),
            fmt(row["attention_gap_lower_bound_ms"]),
            fmt(row["rmsnorm_gap_ms"]),
            fmt(row["swiglu_gap_ms"]),
            fmt(row["remaining_gpu_gap_ms"]),
            fmt(row["identified_gap_share_pct"], 1) + "%",
        ]
        for row in phase_rows
    ]

    markdown = [
        "# AIC 与真机 CUDA Kernel 误差定位（更正后的算子对齐）",
        "",
        "## 结论",
        "",
        "这次定位确认，主要缺口不在已采集的 BF16 GEMM 表本身，而在真机实际执行的纯 Torch 融合算子回退路径。三个真机任务均设置了 `SGLANG_FORCE_FUSED_OP_BACKEND=torch`：RMSNorm 被拆成 FP32 cast、pow、mean、rsqrt、乘法和回写等多个 CUDA kernel；SwiGLU 被拆成 SiLU 与 BF16 乘法。AIC 的 LLAMA/Qwen 模型则把 `add_norm` 和 `act_gate` 各表示为一个解析式 `ElementWise` 项，因此没有复现这些 kernel 边界和 FP32 accumulator/cast 成本。",
        "",
        "另外，上一份算子报告中把 AIC 的整个 `attention` 项与 trace 中的 Triton FMHA 主 kernel 直接比较，边界并不一致：AIC 的 prefill attention 项包含解析式的 RoPE/KV（Qwen3 还包括 QK norm）附加项；decode attention 项则只查询 attention 表、没有同类附加项。本报告将 FMHA 主 kernel 与附加项分开，替代此前该列的解读。",
        "",
        "## AIC 总和与 trace GPU 总和",
        "",
        "`trace GPU sum` 是一次 CUDA Graph replay 中 CUDA kernel/memcpy/memset duration 的串行和，不是 HTTP 或 scheduler 端到端时间；五次真机时延基线仍保留在原单步对比结果中。",
        "",
        "| 模型 | 阶段 | AIC 总 ms | trace GPU sum ms | 偏差 |",
        "| --- | --- | ---: | ---: | ---: |",
    ]
    markdown.extend("| " + " | ".join(row) + " |" for row in aggregate_table)
    markdown += [
        "",
        "## 已确认的纯 Torch ElementWise 回退缺口",
        "",
        "真机列只统计可从 native RMSNorm 和 native `SiluAndMul` 调用链明确识别的 kernel；不把有歧义的 residual add 和通用 copy 强行归类，因此它们是保守下界。负偏差表示 AIC 低估。",
        "",
        "| 模型 | 阶段 | AIC add+norm | trace native RMSNorm | 偏差 | AIC act+gate | trace native SwiGLU | 偏差 |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    markdown.extend("| " + " | ".join(row) + " |" for row in elementwise_table)
    markdown += [
        "",
        "在当前运行路径中，AIC `ElementWise` 只有读写字节、经验带宽系数和固定延迟；它没有读取单独的 norm/activation/rope 校准表。相反，真机的 `RMSNorm.forward_native` 明确先转 FP32，再做归一化和转回 BF16；`SiluAndMul.forward_native` 明确执行 `F.silu(x) * x`。这就是小模型和短步长中非 GEMM 部分被明显低估的直接原因。",
        "",
        "## Attention 边界拆分",
        "",
        "`trace attention auxiliary` 为已识别的 RoPE 形状 BF16 乘法、RoPE cat/gather、QK norm 与 KV store 之和；它不含有歧义的 residual/copy/index，因此是下界。AIC prefill 的 auxiliary 是源码中 `ContextAttention.query()` 的解析式项；AIC decode 的 `GenerationAttention.query()` 不添加该类项，所以为 0。",
        "",
        "| 模型 | 阶段 | AIC FMHA table | trace FMHA | core 偏差 | AIC aux | trace aux 下界 | aux 偏差 | AIC attention 总 | trace attention 下界 | 总偏差 |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    markdown.extend("| " + " | ".join(row) + " |" for row in attention_table)
    markdown += [
        "",
        "这说明 GEMM 和 FMHA 主核的采集覆盖确实有价值，但仅继续加密 GEMM/FMHA shape 数据无法补齐当前总误差。尤其是 decode 路径，AIC 的 generation attention 项没有把 RoPE/QK norm/KV 写入以同样方式并入预测，而 trace 中这些 kernel 是可见的。",
        "",
        "## 总误差的下界归因",
        "",
        "下表以 `trace GPU sum − AIC total` 为观测缺口；各项为 trace−AIC。attention 为已识别辅助 kernel 的下界，`剩余` 包含 residual、泛化 copy、cache/index 和未归类 kernel。GEMM 为负数意味着它抵消了部分低估。",
        "",
        "| 模型 | 阶段 | 观测缺口 | GEMM | attention 下界 | RMSNorm | SwiGLU | 剩余 | 已识别比例 |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    markdown.extend("| " + " | ".join(row) + " |" for row in gap_table)
    markdown += [
        "",
        "## 代码路径证据",
        "",
        "- 真机命令记录在 `models/<model>/command.txt`，均包含 `SGLANG_FORCE_FUSED_OP_BACKEND=torch`；该全局开关把 `BaseFusedOp` 派发到 `forward_native`。",
        "- `python/sglang/srt/layers/layernorm.py` 的 `RMSNorm.forward_native` 使用 FP32 的 `pow → mean → rsqrt → multiply → cast` 链；trace 中对应的 FP32 kernel 与 cast kernel 已单独汇总。",
        "- `python/sglang/srt/layers/activation.py` 的 `SiluAndMul.forward_native` 是 `F.silu(...) * ...`；trace 中同一 grid 的 SiLU 和 BF16 multiply 被配对统计。",
        "- AIC 的 `sdk/models/llama.py` 为 `add_norm`/`act_gate` 建立 `ElementWise`，`sdk/operations/elementwise.py` 通过 `query_mem_op` 解析估时，并非本机采集表。",
        "- AIC 的 `sdk/operations/attention.py`：prefill `ContextAttention.query()` 在 table 外解析加入 RoPE/QK norm/KV 项；decode `GenerationAttention.query()` 仅调用 generation attention table。",
        "",
        "## 最小成本的后续校准方向",
        "",
        "1. 保持当前真实基线的模型、BF16、Triton attention 和 CUDA Graph 不变，新增同 shape 的 RMSNorm、SwiGLU、RoPE 和 KV store 微基准或 trace-derived 校正项；不要用 GEMM 表替代这些项。",
        "2. 若能提供与当前 Torch/CUDA/SM80 ABI 匹配的 `sglang-kernel`，去掉强制 Torch 回退后重新采样；这将改变实际 kernel 边界，因此需要将 GEMM、attention 和新辅助算子一起重新校准。",
        "3. 修改 AIC decode attention 组合前，先用本报告的 `aic_attention_internal_breakdown.csv` 验证附加项是否按同一边界计入；不能再把整项和单个 FMHA kernel 直接对比。",
        "",
        "## 复现信息和产物",
        "",
        f"- 源码 revision：`{metadata.get('source_revision', 'not recorded')}`。",
        f"- AIC 对照基线：`{baseline}`。",
        f"- shape：batch=1，prefill={input_len} token，首次 decode=1 token；真机保留 prefill `tc_piecewise` 与 decode `full` CUDA Graph。",
        "- 数据：`aic_raw_kernel_phase_summary.csv`、`aic_raw_kernel_component_alignment.csv`、`aic_attention_internal_breakdown.csv`、`raw_cuda_kernel_families.csv` 和 `aic_raw_kernel_root_cause.json`。",
        "",
    ]
    markdown_path = analysis / "aic_raw_kernel_root_cause.md"
    markdown_path.write_text("\n".join(markdown), encoding="utf-8")

    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9)
    document.add_heading("AIC 与真机 CUDA Kernel 误差定位（更正后的算子对齐）", 0)
    document.add_paragraph(
        "结论：当前主要缺口来自 SGLANG_FORCE_FUSED_OP_BACKEND=torch 导致的纯 Torch "
        "RMSNorm/SwiGLU 多 kernel 回退，以及 AIC 对这些部分的单一解析 ElementWise 表达；"
        "不是已采集 BF16 GEMM 表本身的主要问题。"
    )
    document.add_heading("AIC 总和与 trace GPU 总和", level=1)
    add_docx_table(document, ["模型", "阶段", "AIC总", "trace GPU总", "偏差"], aggregate_table)
    document.add_paragraph(
        "trace GPU sum 是一次 CUDA Graph replay 的 CUDA event duration 串行和，不等于 HTTP 或 scheduler 端到端时延。"
    )
    document.add_heading("纯 Torch ElementWise 回退", level=1)
    add_docx_table(
        document,
        ["模型", "阶段", "AIC add+norm", "trace RMSNorm", "偏差", "AIC act+gate", "trace SwiGLU", "偏差"],
        elementwise_table,
    )
    document.add_paragraph(
        "真机 RMSNorm 列为可明确识别的 FP32/cast kernel 下界；SwiGLU 列为同一 grid 的 SiLU 和 BF16 乘法之和。"
    )
    document.add_heading("Attention 边界拆分", level=1)
    add_docx_table(
        document,
        ["模型", "阶段", "AIC FMHA", "trace FMHA", "core偏差", "AIC aux", "trace aux下界", "aux偏差", "AIC attn总", "trace attn下界", "总偏差"],
        attention_table,
    )
    document.add_paragraph(
        "这张表更正了此前将 AIC attention 整项和单个 Triton FMHA kernel 直接比较的边界不一致问题。"
    )
    document.add_heading("总误差的下界归因", level=1)
    add_docx_table(
        document,
        ["模型", "阶段", "观测缺口", "GEMM", "attention下界", "RMSNorm", "SwiGLU", "剩余", "已识别比例"],
        gap_table,
    )
    document.add_paragraph(
        "各项为 trace GPU sum 减 AIC total；attention 为保守下界，剩余项含 residual、copy、cache/index 和未归类 kernel。"
    )
    document.add_heading("后续校准方向", level=1)
    document.add_paragraph(
        "保持当前 CUDA Graph 真机基线，单独补齐 RMSNorm、SwiGLU、RoPE 和 KV store 的同 shape 数据。"
        "如果切换到 ABI 匹配的 SM80 sglang-kernel，则必须把所有相关算子重新采样，而不是沿用当前回退路径的结论。"
    )
    document.save(analysis / "aic_raw_kernel_root_cause.docx")


if __name__ == "__main__":
    main()
