"""Format and validate the generated stage DOCX against its Markdown source."""
import hashlib
import importlib.util
import json
import re
import zipfile
from pathlib import Path
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
SIM=Path("/home/ms-admin/workspace1/yiheng/ali/sglang-pr33824-merged/tools/sglang-simulator")
OUT=Path(__file__).resolve().parent
MD=SIM/"docs/pr33824_dense_models_stage_20260929_1906.md"
DOC=MD.with_suffix(".docx")
spec=importlib.util.spec_from_file_location("stage_renderer",SIM/"scripts/render_markdown_report_docx.py")
renderer=importlib.util.module_from_spec(spec); spec.loader.exec_module(renderer)
doc=Document(DOC)
removed=0
for root in [doc._element,doc.styles.element]:
    for fonts in root.iter(qn("w:rFonts")):
        for attr in list(fonts.attrib):
            if etree.QName(attr).localname.lower().endswith("theme"):
                del fonts.attrib[attr]; removed+=1
        for attr in ("ascii","hAnsi","eastAsia","cs"): fonts.set(qn("w:"+attr),"SimSun")
for section in doc.sections:
    for container in (section.header,section.footer):
        for fonts in container._element.iter(qn("w:rFonts")):
            for attr in list(fonts.attrib):
                if etree.QName(attr).localname.lower().endswith("theme"):
                    del fonts.attrib[attr]; removed+=1
            for attr in ("ascii","hAnsi","eastAsia","cs"): fonts.set(qn("w:"+attr),"SimSun")
captions=set(re.findall(r"^!\[([^\]]*)\]\([^)]+\)$",MD.read_text(encoding="utf-8"),flags=re.M))
section=doc.sections[0]
available_width=(section.page_width-section.left_margin-section.right_margin)/914400
for paragraph in doc.paragraphs:
    if paragraph.text in captions: paragraph.paragraph_format.keep_with_next=True
for table in doc.tables:
    table.autofit=False
    n=len(table.columns)
    if n==3: weights=[1.3,4.4,4.4]
    elif "Checkpoint" in table.cell(0,0).text: weights=[2.7,1.7,2.0,1.6,1.3,1.3]
    elif n==5: weights=[1.8,2.8,1.3,2.8,1.3]
    else: weights=[1.8]+[1.35]*(n-1)
    widths=[available_width*x/sum(weights) for x in weights]
    for row in table.rows:
        for cell,width in zip(row.cells,widths): cell.width=Inches(width)
        prop=row._tr.get_or_add_trPr()
        if prop.find(qn("w:cantSplit")) is None:
            no_split=OxmlElement("w:cantSplit"); prop.append(no_split)
    prop=table.rows[0]._tr.get_or_add_trPr()
    if prop.find(qn("w:tblHeader")) is None:
        repeat=OxmlElement("w:tblHeader"); prop.append(repeat)
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs: run.bold=True
renderer.apply_simsun(doc)
doc.save(DOC)
markdown=MD.read_text(encoding="utf-8")
expected=[]; table_count=0; image_files=[]
inside_table=False
for line in markdown.splitlines():
    stripped=line.strip()
    if not stripped: inside_table=False; continue
    if stripped.startswith("|"):
        if not inside_table: table_count+=1; inside_table=True
        if not renderer.is_table_separator(stripped): expected+=renderer.table_cells(stripped)
        continue
    inside_table=False
    image=re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$",stripped)
    if image:
        expected.append(renderer.clean_inline(image[1]))
        image_files.append((MD.parent/image[2]).resolve())
        continue
    heading=re.match(r"^#{1,6}\s+(.+)$",stripped)
    if heading: expected.append(renderer.clean_inline(heading[1]))
    elif stripped.startswith("- "): expected.append(renderer.clean_inline(stripped[2:]))
    else: expected.append(renderer.clean_inline(stripped))
def paragraph_text(p): return "".join(node.text or "" for node in p.iter(qn("w:t")))
actual=[]
for element in doc._element.body:
    if element.tag==qn("w:p"):
        value=paragraph_text(element)
        if value: actual.append(value)
    elif element.tag==qn("w:tbl"):
        for row in element.findall(qn("w:tr")):
            for cell in row.findall(qn("w:tc")):
                actual.append("\n".join(paragraph_text(p) for p in cell.findall(qn("w:p"))))
assert actual==expected, "Markdown/DOCX text or table content mismatch"
assert len(doc.tables)==table_count==11
assert len(doc.inline_shapes)==len(image_files)==6
with zipfile.ZipFile(DOC) as archive:
    image_hashes={hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith("word/media/")}
    assert image_hashes=={hashlib.sha256(path.read_bytes()).hexdigest() for path in image_files}
    for name in ["word/document.xml","word/styles.xml"]+[n for n in archive.namelist() if re.fullmatch(r"word/(header|footer)\d+\.xml",n)]:
        root=etree.fromstring(archive.read(name))
        for fonts in root.iter(qn("w:rFonts")):
            assert all(fonts.get(qn("w:"+attr))=="SimSun" for attr in ("ascii","hAnsi","eastAsia","cs")),(name,dict(fonts.attrib))
            assert not any(etree.QName(attr).localname.lower().endswith("theme") for attr in fonts.attrib)
source=json.loads((OUT/"source_manifest.json").read_text())
for path,digest in source["source_sha256"].items():
    assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==digest,path
summary=json.loads((OUT/"summary.json").read_text())
assert len(summary["models"])==10 and summary["training_timing_rows"]==98127 and summary["heldout_timing_rows"]==164241
result={"markdown":str(MD),"docx":str(DOC),"tables":table_count,"embedded_figures":len(image_files),
        "text_and_table_cells_identical":True,"word_text_font":"SimSun","all_explicit_font_slots_checked":True,
        "removed_theme_font_attributes":removed,"source_files_still_unchanged":len(source["source_sha256"]),
        "md_sha256":hashlib.sha256(MD.read_bytes()).hexdigest(),"docx_sha256":hashlib.sha256(DOC.read_bytes()).hexdigest(),
        "scope":"Dense stage document formatting/content/data validation only. No active experimental implementation or process altered."}
(OUT/"document_validation.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n")
print(json.dumps(result,ensure_ascii=False))
